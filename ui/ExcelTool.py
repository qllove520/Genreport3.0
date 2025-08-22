import sys
import os
from PyQt5.QtWidgets import (
    QApplication, QWidget, QLabel, QLineEdit, QPushButton,
    QVBoxLayout, QHBoxLayout, QFileDialog, QMessageBox
)
from PyQt5.QtCore import Qt
from core.excel_utils import find_row_by_fuzzy_column_value,write_to_target_sheet
from openpyxl import load_workbook
from config.settings import Reference_msg_path,Testreport_path
from core.settings_manager import SettingsManager

class ExcelTool(QWidget):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("项目台账写入工具")
        self.setMinimumSize(700, 500)

        self.settings_manager = SettingsManager()
        # 优先读取配置文件
        loaded_settings = self.settings_manager.load_settings(
            "acceptance_filling",
            default_settings={
                "data_file": Reference_msg_path,
                "template_file": Testreport_path
            }
        )
        self.data_file = loaded_settings.get("data_file", Reference_msg_path)
        self.template_file = loaded_settings.get("template_file", Testreport_path)

        # 新增输入字段变量
        self.input_fields = {}

        self.init_ui()

    def init_ui(self):
        from PyQt5.QtWidgets import QFormLayout

        main_layout = QVBoxLayout(self)
        form_layout = QFormLayout()
        self.setStyleSheet("""
            QWidget {
                font-family: 'Segoe UI', sans-serif;
                font-size: 14px;
            }
            QLabel {
                color: #333;
                font-weight: bold;
            }
            QLineEdit {
                padding: 8px;
                border: 1px solid #ccc;
                border-radius: 5px;
                font-size: 14px;
            }
            QLineEdit:focus {
                border: 1px solid #007bff;
            }
            QPushButton {
                background-color: #007bff;
                color: white;
                border-radius: 5px;
                font-size: 14px;
                padding: 8px 15px;
            }
            QPushButton:hover {
                background-color: #0056b3;
            }
            QPushButton:pressed {
                background-color: #004085;
            }
            QMessageBox {
                background-color: #f0f0f0;
                font-size: 14px;
            }
        """)

        # Helper to create file selection row
        def create_file_selection_row(label_text, line_edit_widget, browse_func):
            file_widget = QWidget()
            file_layout = QHBoxLayout(file_widget)
            file_layout.setContentsMargins(0, 0, 0, 0)
            line_edit_widget.setReadOnly(True)
            line_edit_widget.setStyleSheet("""
                QLineEdit {
                    background-color: #f5f5f5;
                }
            """)
            btn_browse = QPushButton("浏览...")
            btn_browse.clicked.connect(browse_func)
            btn_browse.setStyleSheet("""
                QPushButton {
                    background-color: #2196F3;
                    color: white;
                    border-radius: 5px;
                    font-size: 14px;
                    padding: 8px 15px;
                }
                QPushButton:hover {
                    background-color: #1976D2;
                }
                QPushButton:pressed {
                    background-color: #1565C0;
                }
            """)
            file_layout.addWidget(line_edit_widget)
            file_layout.addWidget(btn_browse)
            form_layout.addRow(label_text, file_widget)

        # 数据台账
        self.label_data = QLabel("未选择 数据台账 文件")
        self.data_file_display = QLineEdit(f"{os.path.join(self.data_file)}")
        create_file_selection_row("项目台账（读取）:", self.data_file_display, self.choose_data_file)

        # 写入模板
        self.label_template = QLabel("未选择 写入模板 文件")
        self.template_file_display = QLineEdit(f"{os.path.join(self.template_file)}")
        create_file_selection_row("验收测试报告（写入）:", self.template_file_display, self.choose_template_file)

        # 关键词
        self.input_keyword = QLineEdit()
        self.input_keyword.setPlaceholderText("请输入唯一任务书编号，例如：CPKFLX20230619001")
        form_layout.addRow("任务书编号：:", self.input_keyword)

        # 额外字段
        extra_fields = [
            "测试单号", "申请理由", "开始时间", "结束时间", "测试依据", "测试范围"
        ]

        extra_tips =["CPKFLX20240426001","新产品导入","2025/07/14","2025/07/21","窗口式照相机技术需求规格书V1.3.xlsx","回归上一轮Bug、按照整机验收标准：底层软件、上层应用、机电安全"]

        for i, field in enumerate(extra_fields):
            input_box = QLineEdit()
            input_box.setPlaceholderText(f"可选填写 例如： {extra_tips[i]}")
            input_box.setStyleSheet("""
                QLineEdit {
                    background-color: #e9ecef;
                }
            """)
            form_layout.addRow(field + ":", input_box)
            self.input_fields[field] = input_box

        main_layout.addLayout(form_layout)

        # 提交按钮
        btn_process = QPushButton("提取并写入")
        btn_process.clicked.connect(self.process)
        btn_process.setFixedHeight(40)
        btn_process.setStyleSheet("""
            QPushButton {
                background-color: #007bff;
                color: white;
                border-radius: 8px;
                font-size: 16px;
                font-weight: bold;
                padding: 10px 20px;
            }
            QPushButton:hover {
                background-color: #0056b3;
            }
            QPushButton:pressed {
                background-color: #004085;
            }
        """)
        main_layout.addWidget(btn_process, alignment=Qt.AlignCenter)

        self.setLayout(main_layout)

    def choose_data_file(self):
        initial_dir = os.path.dirname(self.data_file) if self.data_file else ""
        path, _ = QFileDialog.getOpenFileName(self, "选择 数据台账 文件", initial_dir, "Excel Files (*.xlsx *.xlsm)")
        if path:
            self.data_file = path
            self.data_file_display.setText(path)
            self.label_data.setText(f"已选择数据台账: {path}")
            # 保存到配置
            self.settings_manager.save_settings(
                "acceptance_filling",
                {"data_file": self.data_file, "template_file": self.template_file}
            )

    def choose_template_file(self):
        initial_dir = os.path.dirname(self.template_file) if self.template_file else ""
        path, _ = QFileDialog.getOpenFileName(self, "选择 写入模板 文件", initial_dir, "Excel Files (*.xlsx *.xlsm)")
        if path:
            self.template_file = path
            self.template_file_display.setText(path)
            self.label_template.setText(f"已选择模板文件: {path}")
            # 保存到配置
            self.settings_manager.save_settings(
                "acceptance_filling",
                {"data_file": self.data_file, "template_file": self.template_file}
            )

    def process(self):
        if not self.data_file or not self.template_file:
            QMessageBox.warning(self, "错误", "请先选择 数据台账 和 写入模板")
            return

        keyword = self.input_keyword.text().strip()
        if not keyword:
            QMessageBox.warning(self, "错误", "请输入关键词")
            return

        try:
            # 提取主数据
            target_fields = ['项目编号', '项目名称', '内部型号', '产品名称', '项目经理', '产品经理', '负责人']
            cell_mapping = {
                '项目编号': 'D2',
                '项目名称': 'H2',
                '项目经理': 'U2',
                '内部型号': 'D3',
                '产品名称': 'H3',
                '产品经理': 'U3',
                '负责人': 'U4'
            }

            result = find_row_by_fuzzy_column_value(
                file_path=self.data_file,
                key_column='项目_产品',
                key_value=keyword,
                target_columns=target_fields
            )

            if not result:
                QMessageBox.warning(self, "未找到", "台账中未找到匹配行")
                return

            # 提取额外输入内容
            extra_data = {
                field: self.input_fields[field].text().strip()
                for field in self.input_fields
                if self.input_fields[field].text().strip()
            }

            extra_cell_mapping = {
                '测试单号': 'O2',
                '申请理由': 'D4',
                '开始时间': 'H4',
                '结束时间': 'O4',
                '测试依据': 'E6',
                '测试范围': 'E7'
            }

            # 合并数据并写入
            wb = load_workbook(self.template_file,keep_vba=True)
            sheet_name = '验收测试结果'
            if sheet_name not in wb.sheetnames:
                raise ValueError(f"写入模板缺少工作表：{sheet_name}")
            sheet = wb[sheet_name]

            # 主数据写入
            for key, cell in cell_mapping.items():
                sheet[cell] = result.get(key, "")

            # 附加字段写入（仅填写的才写）
            for key, cell in extra_cell_mapping.items():
                if key in extra_data:
                    sheet[cell] = extra_data[key]

            wb.save(self.template_file)

            QMessageBox.information(self, "成功", "数据已成功写入 Excel 模板！")

        except Exception as e:
            QMessageBox.critical(self, "错误", f"发生错误：\n{str(e)}")
